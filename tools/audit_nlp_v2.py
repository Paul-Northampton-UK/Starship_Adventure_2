"""Audit NLP v2 capability registry and export summaries.

Usage:
  python -m tools.audit_nlp_v2 --format text
  python -m tools.audit_nlp_v2 --format markdown --out docs/nlp_v2_capabilities.md
  python -m tools.audit_nlp_v2 --format csv --out docs/nlp_v2_capabilities.csv
  python -m tools.audit_nlp_v2 --format xlsx --out docs/nlp_v2_capabilities.xlsx
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List

from engine.nlp_v2.parser import INTENT_GROUPS, NLPCommandParserV2

PLACEHOLDER_DEFAULTS = {
    "movement verb": "go",
    "direction": "north",
    "verb": "use",
    "object": "cabinet",
    "target": "door",
    "npc": "guard",
    "item": "keycard",
    "item_a": "herb",
    "item_b": "solvent",
    "resource": "berries",
    "trap": "trap",
    "spell": "fireball",
    "power": "mana",
    "entity": "spirit wolf",
    "effect": "flame",
    "from": "lead",
    "to": "gold",
    "name": "moon rite",
    "place": "airlock",
    "weapon": "pistol",
    "ammo": "cell",
    "container": "locker",
    "preposition": "in",
    "scope": "room",
}

DEFAULT_EXT = {
    "text": "txt",
    "markdown": "md",
    "json": "json",
    "csv": "csv",
    "xlsx": "xlsx",
}

HEADERS = ["Group", "Intent", "Verbs", "Form", "Slots", "Example"]


def _summaries(capabilities: Dict[str, Dict[str, List[str]]]) -> dict:
    total_intents = len(capabilities)
    total_forms = sum(len(entry["forms"]) for entry in capabilities.values())
    slot_names = {
        slot for entry in capabilities.values() for slot in entry.get("slots", [])
    }
    return {
        "total_intents": total_intents,
        "total_forms": total_forms,
        "distinct_slots": len(slot_names),
    }


def _clean_whitespace(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _auto_example(form: str) -> str:
    example = form
    for placeholder, sample in PLACEHOLDER_DEFAULTS.items():
        example = example.replace(f"<{placeholder}>", sample)
    example = re.sub(r"<[^>]+>", "", example)
    return _clean_whitespace(example.replace("/", "/"))


def _build_rows(capabilities: Dict[str, Dict[str, List[str]]], include_examples: bool) -> List[dict]:
    rows: list[dict] = []
    for intent in sorted(capabilities):
        entry = capabilities[intent]
        group = INTENT_GROUPS.get(intent, "Other")
        verbs = ", ".join(entry["verbs"]) if entry["verbs"] else "-"
        slots = ", ".join(entry["slots"]) if entry["slots"] else ""
        forms = entry["forms"] or ["-"]
        examples = entry.get("examples", []) if include_examples else []
        for idx, form in enumerate(forms):
            example = ""
            if include_examples:
                example = examples[idx] if idx < len(examples) else _auto_example(form)
            rows.append(
                {
                    "Group": group,
                    "Intent": intent,
                    "Verbs": verbs,
                    "Form": form,
                    "Slots": slots,
                    "Example": example,
                }
            )
    return rows


def _format_text(rows: List[dict]) -> str:
    grouped: dict[str, dict[str, list[dict]]] = defaultdict(lambda: defaultdict(list))
    for row in rows:
        grouped[row["Group"]][row["Intent"]].append(row)

    lines: list[str] = []
    for group in sorted(grouped):
        lines.append(f"{group}:")
        intents = grouped[group]
        for intent in sorted(intents):
            lines.append(f"  {intent}")
            intent_rows = intents[intent]
            verbs = intent_rows[0]["Verbs"] or "-"
            slots = intent_rows[0]["Slots"] or "(none)"
            lines.append(f"    Verbs: {verbs}")
            lines.append(f"    Slots: {slots}")
            lines.append("    Forms:")
            for row in intent_rows:
                form_line = f"      - {row['Form']}"
                if row["Example"]:
                    form_line += f" (e.g., \"{row['Example']}\")"
                lines.append(form_line)
        lines.append("")
    return "\n".join(lines).rstrip()


def _format_markdown(rows: List[dict]) -> str:
    lines = ["| Group | Intent | Verbs | Form | Slots | Example |", "| --- | --- | --- | --- | --- | --- |"]
    for row in rows:
        verbs = row["Verbs"].replace(", ", "<br>") or "-"
        slots = row["Slots"].replace(", ", "<br>") if row["Slots"] else "(none)"
        example = row["Example"] or ""
        lines.append(
            f"| {row['Group']} | {row['Intent']} | {verbs} | {row['Form']} | {slots} | {example} |"
        )
    return "\n".join(lines)


def _format_json(capabilities: Dict[str, Dict[str, List[str]]]) -> str:
    enriched: dict[str, dict] = {}
    for intent, entry in capabilities.items():
        enriched[intent] = {
            **entry,
            "group": INTENT_GROUPS.get(intent, "Other"),
        }
    return json.dumps(enriched, indent=2)


def _format_summary(capabilities: Dict[str, Dict[str, List[str]]], fmt: str) -> str:
    summary = _summaries(capabilities)
    if fmt == "markdown":
        return "\n".join(
            [
                "| Metric | Count |",
                "| --- | --- |",
                f"| Total intents | {summary['total_intents']} |",
                f"| Total forms | {summary['total_forms']} |",
                f"| Distinct slots | {summary['distinct_slots']} |",
            ]
        )
    if fmt == "json":
        return json.dumps({"summary": summary}, indent=2)
    return "\n".join(
        [
            f"Total intents: {summary['total_intents']}",
            f"Total forms: {summary['total_forms']}",
            f"Distinct slots: {summary['distinct_slots']}",
        ]
    )


def _write_csv(path: Path, rows: List[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def _write_xlsx(path: Path, rows: List[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NLP v2"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row[h] or "" for h in HEADERS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"
    for idx, header in enumerate(HEADERS, start=1):
        max_len = max(
            [len(header)]
            + [len(str(row.get(header, "") or "")) for row in rows]
        )
        width = max_len + 2
        sheet.column_dimensions[get_column_letter(idx)].width = width
    workbook.save(path)


def _resolve_output_path(raw_path: str, fmt: str) -> Path:
    path = Path(raw_path)
    if raw_path.endswith(("/", "\\")) or path.is_dir():
        return path / f"nlp_v2_capabilities.{DEFAULT_EXT[fmt]}"
    if not path.suffix:
        return path.with_suffix(f".{DEFAULT_EXT[fmt]}")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Audit NLP v2 capabilities (verbs, forms, slots)."
    )
    parser.add_argument(
        "--format",
        choices=["text", "markdown", "json", "csv", "xlsx"],
        default="text",
        help="Output format.",
    )
    parser.add_argument("--out", type=str, help="Optional output path/directoy.")
    parser.add_argument(
        "--summary",
        action="store_true",
        help="Only output summary counts (text/markdown/json formats).",
    )
    parser.add_argument(
        "--examples",
        dest="examples",
        action="store_true",
        default=True,
        help="Include auto-generated examples (default).",
    )
    parser.add_argument(
        "--no-examples",
        dest="examples",
        action="store_false",
        help="Omit example column/content.",
    )
    args = parser.parse_args()

    capabilities = NLPCommandParserV2.get_capabilities()
    rows = _build_rows(capabilities, include_examples=args.examples)
    fmt = args.format

    output_path: Path | None = None
    if args.out:
        output_path = _resolve_output_path(args.out, fmt)

    if args.summary and fmt in {"csv", "xlsx"}:
        raise SystemExit("--summary is only supported for text/markdown/json outputs.")

    if args.summary:
        content = _format_summary(capabilities, fmt)
        if output_path:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(content, encoding="utf-8")
        else:
            print(content)
        return

    if fmt == "markdown":
        content = _format_markdown(rows)
    elif fmt == "json":
        content = _format_json(capabilities)
    elif fmt == "text":
        content = _format_text(rows)
    elif fmt == "csv":
        if not output_path:
            raise SystemExit("--out is required for csv format.")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        _write_csv(output_path, rows)
        return
    elif fmt == "xlsx":
        if not output_path:
            raise SystemExit("--out is required for xlsx format.")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        _write_xlsx(output_path, rows)
        return
    else:
        raise SystemExit(f"Unsupported format: {fmt}")

    if output_path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(content, encoding="utf-8")
    else:
        print(content)


if __name__ == "__main__":
    main()
